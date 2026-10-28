"""SLA source grounding, chosen assignments, edits, and downloadable exports."""

import io

import pytest
from api.index import app
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pypdf import PdfReader

from talent_estimation.catalog import load_default_catalog
from talent_estimation.sla import (
    PLACEHOLDER,
    SlaDraft,
    SlaDraftRequest,
    SlaRow,
    SlaSeed,
    export_sla_pdf,
    export_sla_word,
    make_sla,
    read_assignment_plan,
    seed_from_staffing,
)
from talent_estimation.staffing import Assignment, Demand, StaffingRequest, Talent, TalentProject


def seed():
    return SlaSeed(
        project="Project Example",
        rows=[SlaRow(role="DE", talent="Selected person", fte=0.35)],
        tasks={"DE": ["Build ingestion pipeline"]},
    )


def test_only_human_selected_talent_and_exact_fte_are_used():
    body = StaffingRequest(
        project=TalentProject(demands=[Demand(role="DE", fte=1)]),
        talents=[Talent(id="chosen", name="Chosen"), Talent(id="other", name="Not selected")],
        assignments=[Assignment(talent_id="chosen", role="DE", fte=0.35)],
    )
    result = seed_from_staffing(body)
    assert [(r.talent, r.fte) for r in result.rows] == [("Chosen", 0.35)]


def test_import_plan_never_uses_recommendation_rows():
    w = Workbook()
    s = w.active
    s.title = "Assignments"
    s.append(["Project", "Role", "Talent", "Assigned FTE", "Status"])
    s.append(["P", "Engineer", "Chosen", "0,5", "Over capacity"])
    w.create_sheet("Recommendations").append(["Someone else"])
    b = io.BytesIO()
    w.save(b)
    result = read_assignment_plan(b.getvalue())
    assert len(result.rows) == 1 and result.rows[0].role == "AI"
    assert result.rows[0].fte == 0.5 and result.warnings


def test_missing_project_context_requires_document_or_manual_context():
    with pytest.raises(ValueError, match="KAK/project charter"):
        make_sla(SlaDraftRequest(seed=seed()), [], load_default_catalog())


def test_grounded_content_catalog_scope_and_missing_facts():
    catalog = load_default_catalog()
    catalog.roles["DE"].in_scope = "Custom pipeline scope"
    catalog.roles["DE"].out_of_scope = "Custom exclusion"
    body = SlaDraftRequest(seed=seed(), project_info={"start": "2027-01-01", "end": "2027-03-31"})
    d = make_sla(
        body,
        [
            (
                "Charter.txt",
                "Background\nA documented project background for this engagement.\nObjectives\nDeliver a reviewed pipeline.\nProject Owner / PM: Owner Example",
            )
        ],
        catalog,
    )
    assert "documented project background" in d.background
    assert d.owner_pm == "Owner Example" and d.chapter == PLACEHOLDER
    assert d.period == "2027-01-01 s.d. 2027-03-31"
    assert "Custom pipeline scope" in d.rows[0].in_scope
    assert "Build ingestion pipeline" in d.rows[0].in_scope
    assert d.rows[0].out_scope == "Custom exclusion"
    assert d.sources == ["Charter.txt"]
    assert "4-6 jam kerja" in d.communication and "2 jam kerja" in d.communication


def test_exports_retain_reviewers_text_and_table_edits():
    d = SlaDraft(
        project_name="Edited project",
        intake="Edited intake <safe> & retained",
        rows=[
            SlaRow(
                role="Edited role",
                talent="Edited name",
                fte=0.45,
                in_scope="Edited deliverable",
                out_scope="Edited exclusion",
            )
        ],
    )
    doc = Document(io.BytesIO(export_sla_word(d)))
    assert "Edited intake <safe> & retained" in "\n".join(p.text for p in doc.paragraphs)
    assert doc.tables[0].cell(1, 2).text == "Edited name"
    assert doc.tables[0].cell(1, 3).text == "0.45"
    pdf = PdfReader(io.BytesIO(export_sla_pdf(d)))
    text = "\n".join(p.extract_text() for p in pdf.pages)
    assert "Edited intake <safe> & retained" in text
    assert "Edited name" in text and "Edited deliverable" in text


def test_pdf_long_table_wraps_and_continues_across_pages():
    d = SlaDraft(
        rows=[
            SlaRow(
                talent=f"Talent {i}",
                fte=0.5,
                in_scope="Deliverable details " * 60,
                out_scope="Excluded work " * 35,
            )
            for i in range(15)
        ]
    )
    pdf = PdfReader(io.BytesIO(export_sla_pdf(d)))
    text = "\n".join(p.extract_text() for p in pdf.pages)
    assert len(pdf.pages) > 2 and "Talent 14" in text


def test_api_draft_upload_validation_and_both_exports():
    client = TestClient(app)
    body = SlaDraftRequest(seed=seed())
    bad = client.post("/api/py/sla/draft", data={"settings": body.model_dump_json()})
    assert bad.status_code == 400
    response = client.post(
        "/api/py/sla/draft",
        data={"settings": body.model_dump_json()},
        files={
            "documents": (
                "Charter.txt",
                b"Background\nThe project delivers a data platform for business review and analysis.",
            )
        },
    )
    assert response.status_code == 200, response.text
    draft = response.json()
    draft["rows"].append({"role": "DA", "talent": "Reviewer added", "fte": 0.1})
    assert client.post("/api/py/sla/validate", json=draft).status_code == 200
    assert client.post("/api/py/sla/export/docx", json=draft).content.startswith(b"PK")
    assert client.post("/api/py/sla/export/pdf", json=draft).content.startswith(b"%PDF")
    draft["rows"][0]["fte"] = -1
    assert client.post("/api/py/sla/validate", json=draft).status_code == 422
