"""Auditable, role-independent talent ranking and human assignment planning."""

from __future__ import annotations

import hashlib
import io
import math
import re
from collections import defaultdict
from typing import Any, Literal

from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field, model_validator

from talent_estimation.catalog import Catalog

ROLE_NAMES = {"DE": "Data Engineer", "DA": "Data Analyst", "DS": "Data Scientist", "AI": "Engineer"}


def norm(value: Any) -> str:
    """Normalize labels without guessing semantic equivalence."""
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").replace("\ufffd", " ").lower()).strip()


def role_code(value: Any) -> str:
    """Resolve the four data roles, including the survey's Engineer label."""
    text = norm(value)
    for code, name in ROLE_NAMES.items():
        if text in {norm(code), norm(name), "ai engineer" if code == "AI" else norm(name)}:
            return code
    return str(value or "").strip()


def skill_key(role: str, pillar: str, skill: str) -> str:
    """Keep identical names in different role/pillar contexts distinct."""
    return "|".join([role_code(role), norm(pillar), norm(skill)])


class Talent(BaseModel):
    """One respondent; roles are display metadata, never ranking filters."""

    id: str
    name: str
    primary_role: str = ""
    preferred_roles: str = ""
    project: str = ""
    current_fte: float | None = Field(default=None, ge=0, le=10, allow_inf_nan=False)
    capacity_fte: float = Field(default=1, gt=0, le=10, allow_inf_nan=False)
    skills: dict[str, int] = Field(default_factory=dict)

    @model_validator(mode="after")
    def valid_scores(self) -> Talent:
        """Reject invalid self-assessment values instead of improving them silently."""
        if any(v < 1 or v > 5 for v in self.skills.values()):
            raise ValueError("Proficiency must be an integer from 1 to 5.")
        return self


class Requirement(BaseModel):
    """One role-skill requirement on one WBS item."""

    item_id: str
    item_name: str = ""
    role: str
    pillar: str
    skill: str
    classification: Literal["Simple", "Average", "Complex"] = "Average"
    weight: float = Field(default=10, gt=0, le=1e6, allow_inf_nan=False)
    target: int = Field(default=3, ge=2, le=5)
    reason: str = ""


class Demand(BaseModel):
    """Project need from the selected estimation method."""

    role: str
    fte: float | None = Field(default=None, ge=0, le=1000, allow_inf_nan=False)


class TalentProject(BaseModel):
    """Project inputs retained in browser memory with their provenance."""

    name: str = "Project"
    source: str = ""
    method: str = "bottom_up_peak"
    requirements: list[Requirement] = Field(default_factory=list, max_length=10000)
    demands: list[Demand] = Field(default_factory=list, max_length=100)
    warnings: list[str] = Field(default_factory=list)


class Assignment(BaseModel):
    """A human's proposed allocation, not an automatically filled role."""

    talent_id: str
    role: str
    fte: float = Field(gt=0, le=10, allow_inf_nan=False)
    reason: str = ""


class StaffingRequest(BaseModel):
    """Complete stateless scoring request."""

    project: TalentProject
    talents: list[Talent] = Field(max_length=1000)
    assignments: list[Assignment] = Field(default_factory=list, max_length=1000)

    @model_validator(mode="after")
    def unique_ids(self) -> StaffingRequest:
        """Prevent silent allocation to duplicate or unknown respondents."""
        ids = [t.id for t in self.talents]
        if len(ids) != len(set(ids)):
            raise ValueError("Talent IDs must be unique.")
        if any(a.talent_id not in ids for a in self.assignments):
            raise ValueError("Assignment references an unknown talent.")
        roles = {r.role for r in self.project.demands} | {r.role for r in self.project.requirements}
        if any(a.role not in roles for a in self.assignments):
            raise ValueError("Assignment references an unknown project role.")
        pairs = [(a.role, a.talent_id) for a in self.assignments]
        if len(set(pairs)) != len(pairs):
            raise ValueError("Duplicate assignment for the same talent and role.")
        return self


def read_talents(content: bytes, catalog: Catalog) -> dict[str, Any]:
    """Read the two-header breakdown workbook, with optional Current FTE column."""
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    name = "Talent Capability Breakdown"
    if name not in workbook.sheetnames:
        raise ValueError(f"Sheet '{name}' is required.")
    sheet = workbook[name]
    if (sheet.max_row or 0) > 1002 or (sheet.max_column or 0) > 1000:
        raise ValueError("Workbook exceeds 1,000 talents or 1,000 columns.")
    rows = list(sheet.values)
    if len(rows) > 1002 or any(len(r) > 1000 for r in rows):
        raise ValueError("Workbook exceeds 1,000 talents or 1,000 columns.")
    if len(rows) < 3:
        raise ValueError("The workbook has no talent rows.")
    metadata: dict[str, int] = {}
    skill_columns: dict[int, str] = {}
    warnings: list[str] = []
    current_role = ""
    for index, (top, label) in enumerate(zip(rows[0], rows[1], strict=True)):
        header = norm(top) or norm(label)
        if header in {
            "nama",
            "name",
            "role utama",
            "primary role",
            "preferred role",
            "tribe project",
            "current fte",
            "fte saat ini",
            "capacity fte",
        }:
            metadata[header] = index
            continue
        if top:
            current_role = role_code(top)
        if not label or current_role not in catalog.roles:
            continue
        matches = [
            (p, s)
            for p, skills in catalog.roles[current_role].pillars.items()
            for s in skills
            if norm(s) == norm(label)
        ]
        if len(matches) != 1:
            warnings.append(f"Column {index + 1}: '{label}' has no unique catalog match; excluded.")
            continue
        pillar, skill = matches[0]
        key = skill_key(current_role, pillar, skill)
        if key in skill_columns.values():
            raise ValueError(f"Duplicate skill column: {label}.")
        skill_columns[index] = key
    name_index = metadata.get("nama", metadata.get("name"))
    if name_index is None:
        raise ValueError("The first header must include Nama or Name.")
    if not skill_columns:
        raise ValueError("No skill columns match the active catalog.")

    def value(row: tuple[Any, ...], *labels: str) -> Any:
        return next((row[metadata[label]] for label in labels if label in metadata), None)

    talents: list[Talent] = []
    names: set[str] = set()
    for row_number, row in enumerate(rows[2:], 3):
        name_value = str(row[name_index] or "").strip()
        if not name_value:
            continue
        if norm(name_value) in names:
            raise ValueError(
                f"Duplicate talent name '{name_value}'; add a unique identifier to the name."
            )
        names.add(norm(name_value))
        scores: dict[str, int] = {}
        for i, key in skill_columns.items():
            v = row[i]
            if v in (None, ""):
                continue
            try:
                numeric = float(v)
            except (TypeError, ValueError) as exc:
                raise ValueError(
                    f"Row {row_number}, column {i + 1}: invalid proficiency '{v}'."
                ) from exc
            if not math.isfinite(numeric) or numeric != int(numeric) or not 1 <= numeric <= 5:
                raise ValueError(f"Row {row_number}, column {i + 1}: proficiency must be 1–5.")
            scores[key] = int(numeric)
        current = value(row, "current fte", "fte saat ini")
        capacity = value(row, "capacity fte")
        talents.append(
            Talent(
                id=hashlib.sha256(norm(name_value).encode()).hexdigest()[:16],
                name=name_value,
                primary_role=ROLE_NAMES.get(
                    role_code(value(row, "role utama", "primary role")),
                    str(value(row, "role utama", "primary role") or ""),
                ),
                preferred_roles=str(value(row, "preferred role") or ""),
                project=str(value(row, "tribe project") or ""),
                current_fte=float(current) if current not in (None, "") else None,
                capacity_fte=float(capacity) if capacity not in (None, "") else 1,
                skills=scores,
            )
        )
    if not talents:
        raise ValueError("No named talents found.")
    unknown = sum(t.current_fte is None for t in talents)
    if unknown:
        warnings.append(
            f"{unknown} talents have no Current FTE. Enter workload before confirming assignments. Capacity defaults to 1.0 FTE and is editable."
        )
    return {
        "talents": [t.model_dump() for t in talents],
        "warnings": warnings,
        "matched_skills": len(skill_columns),
    }


def project_from_run(run: Any, method: str = "bottom_up_peak") -> TalentProject:
    """Continue directly from a completed, human-editable estimator run."""
    attr = {
        "bottom_up_peak": "bu_peak_fte",
        "ucp_peak": "ucp_peak_fte",
        "bottom_up_average": "bu_fte",
        "ucp_average": "ucp_fte",
    }
    if method not in attr:
        raise ValueError("Unknown FTE estimation method.")
    requirements = []
    for row in run.breakdown:
        weight = run.parameters.use_case_weights[row.classification] * row.quantity
        if str(row.involvement) == "Pendukung":
            weight *= run.parameters.support_factor
        for cap in row.capabilities:
            requirements.append(
                Requirement(
                    item_id=row.item_id,
                    item_name=row.item_name,
                    role=row.role,
                    pillar=cap.pillar,
                    skill=cap.skill,
                    classification=row.classification,
                    weight=weight,
                    reason=cap.reason,
                )
            )
    demands = [Demand(role=e.role, fte=getattr(e, attr[method])) for e in run.estimates]
    return TalentProject(
        name=run.project.name or "Project",
        source="Current estimator run",
        method=method,
        requirements=requirements,
        demands=demands,
    )


def _table(workbook: Any, name: str, required: set[str]) -> list[dict[str, Any]]:
    if name not in workbook.sheetnames:
        raise ValueError(f"Sheet '{name}' is required.")
    rows = list(workbook[name].values)
    for index, row in enumerate(rows[:25]):
        headers = [norm(c) for c in row]
        if required <= set(headers):
            result = []
            for r in rows[index + 1 :]:
                if not any(v is not None and v != "" for v in r):
                    break
                result.append(dict(zip(headers, r, strict=True)))
            return result
    raise ValueError(f"Required headers not found in {name}: {', '.join(sorted(required))}.")


def read_project(content: bytes) -> TalentProject:
    """Import Capability, Breakdown_WBS and cached FTE estimates from a validation file."""
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    if any((s.max_row or 0) > 20000 or (s.max_column or 0) > 1000 for s in workbook):
        raise ValueError("Project workbook is too large.")
    capabilities = _table(workbook, "Capability", {"wbs id", "role", "pilar", "skill"})
    breakdown = _table(workbook, "Breakdown_WBS", {"wbs id", "role", "klasifikasi ucp"})
    lookup = {
        (str(r["wbs id"]), role_code(r["role"])): r
        for r in breakdown
        if r.get("wbs id") and r.get("role")
    }
    project = TalentProject(
        name="Imported validation", source="Validation workbook", method="bottom_up_peak"
    )
    seen = set()
    for row in capabilities:
        if not row.get("wbs id") or not row.get("skill"):
            continue
        role = role_code(row.get("role"))
        item_id = str(row["wbs id"])
        b = lookup.get((item_id, role))
        if b is None:
            project.warnings.append(
                f"{item_id}/{role}: no matching WBS classification; requirement excluded. Add it manually."
            )
            continue
        classification = str(b.get("klasifikasi ucp") or "")
        if classification not in {"Simple", "Average", "Complex"}:
            project.warnings.append(
                f"{item_id}/{role}: invalid UCP class; requirement excluded. Add it manually."
            )
            continue
        key = (item_id, role, norm(row.get("pilar")), norm(row["skill"]))
        if key in seen:
            continue
        seen.add(key)
        effective = b.get("bobot efektif ucp")
        fallback = {"Simple": 5, "Average": 10, "Complex": 15}[classification] * float(
            b.get("qty") or 1
        )
        if b.get("peran") == "Pendukung":
            fallback *= 0.25
        project.requirements.append(
            Requirement(
                item_id=item_id,
                item_name=str(row.get("item wbs") or ""),
                role=role,
                pillar=str(row.get("pilar") or ""),
                skill=str(row["skill"]),
                classification=classification,
                weight=float(effective)
                if effective is not None and float(effective) > 0
                else fallback,
                reason=str(row.get("bagian pekerjaan") or ""),
            )
        )
    try:
        summary = _table(workbook, "Ringkasan", {"kode", "bottom up fte puncak"})
        for row in summary:
            code = str(row.get("kode") or "")
            if code not in {r.role for r in project.requirements}:
                continue
            value = row.get("bottom up fte puncak")
            project.demands.append(
                Demand(role=code, fte=float(value) if isinstance(value, (int, float)) else None)
            )
    except ValueError:
        project.warnings.append("FTE summary unavailable; enter demand per role manually.")
    for role in dict.fromkeys(r.role for r in project.requirements):
        if not any(d.role == role for d in project.demands):
            project.demands.append(Demand(role=role))
    if not project.requirements:
        raise ValueError("No classified capability requirements found in the workbook.")
    return project


def recommend(body: StaffingRequest) -> dict[str, Any]:
    """Rank every talent on project skills, target coverage, and shared capacity."""
    grouped: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    # Repeated skills accumulate effort weight across WBS items, but exact repeated rows do not.
    seen = set()
    for r in body.project.requirements:
        identity = (r.item_id, r.role, norm(r.pillar), norm(r.skill))
        if identity in seen:
            raise ValueError(f"Duplicate requirement: {r.item_id}/{r.role}/{r.skill}.")
        seen.add(identity)
        key = skill_key(r.role, r.pillar, r.skill)
        g = grouped[r.role].setdefault(
            key, {"skill": r.skill, "pillar": r.pillar, "weight": 0.0, "target": r.target}
        )
        g["weight"] += r.weight
        g["target"] = max(g["target"], r.target)
    allocated: dict[str, float] = defaultdict(float)
    for assignment in body.assignments:
        allocated[assignment.talent_id] += assignment.fte
    roles = list(dict.fromkeys([d.role for d in body.project.demands] + list(grouped)))
    result: list[dict[str, Any]] = []
    for role in roles:
        demand = next((d.fte for d in body.project.demands if d.role == role), None)
        assigned = sum(a.fte for a in body.assignments if a.role == role)
        candidates: list[dict[str, Any]] = []
        requirements = grouped.get(role, {})
        total_weight = sum(v["weight"] for v in requirements.values())
        for talent in body.talents:
            detail = []
            for key, requirement in requirements.items():
                level = talent.skills.get(key)
                detail.append(
                    {
                        **requirement,
                        "level": level,
                        "matched": level is not None,
                        "meets_target": level is not None and level >= requirement["target"],
                    }
                )
            proficiency = (
                sum(
                    r["weight"] * ((r["level"] - 1) / 4 if r["level"] is not None else 0)
                    for r in detail
                )
                / total_weight
                if total_weight
                else 0
            )
            coverage = (
                sum(r["weight"] for r in detail if r["meets_target"]) / total_weight
                if total_weight
                else 0
            )
            available = (
                None
                if talent.current_fte is None
                else max(0.0, talent.capacity_fte - talent.current_fte - allocated[talent.id])
            )
            # Rank for the remaining role need, up to one person's capacity. No implicit demand.
            slot = (
                min(max(0.0, demand - assigned), talent.capacity_fte)
                if demand is not None
                else None
            )
            capacity = (
                min(1.0, available / slot) if available is not None and slot and slot > 0 else 0.0
            )
            score = 70 * proficiency + 20 * coverage + 10 * capacity if total_weight else None
            candidates.append(
                {
                    "talent_id": talent.id,
                    "name": talent.name,
                    "primary_role": talent.primary_role,
                    "preferred_roles": talent.preferred_roles,
                    "project": talent.project,
                    "current_fte": talent.current_fte,
                    "capacity_fte": talent.capacity_fte,
                    "allocated_fte": round(allocated[talent.id], 6),
                    "available_fte": available,
                    "score": round(score, 4) if score is not None else None,
                    "skill_score": round(70 * proficiency, 4),
                    "coverage_score": round(20 * coverage, 4),
                    "capacity_score": round(10 * capacity, 4),
                    "coverage_percent": round(100 * coverage, 2),
                    "provisional": talent.current_fte is None or demand is None,
                    "details": detail,
                }
            )
        candidates.sort(
            key=lambda c: (-(c["score"] or 0), -c["coverage_percent"], c["name"].casefold())
        )
        result.append(
            {
                "role": role,
                "required_fte": demand,
                "assigned_fte": assigned,
                "remaining_fte": max(0.0, demand - assigned) if demand is not None else None,
                "candidates": candidates,
                "skills": len(requirements),
            }
        )
    alerts = []
    for t in body.talents:
        if allocated[t.id] == 0:
            continue
        if t.current_fte is None:
            alerts.append(f"{t.name}: Current FTE is unknown; assignment remains provisional.")
        elif t.current_fte + allocated[t.id] > t.capacity_fte + 1e-9:
            alerts.append(
                f"{t.name}: total workload {t.current_fte + allocated[t.id]:.2f} exceeds capacity {t.capacity_fte:.2f} FTE."
            )
    for role_result in result:
        if role_result["remaining_fte"] is None or role_result["remaining_fte"] > 1e-9:
            alerts.append(f"{role_result['role']}: demand is unknown or not fully assigned.")
        if (
            role_result["required_fte"] is not None
            and role_result["assigned_fte"] > role_result["required_fte"] + 1e-9
        ):
            alerts.append(f"{role_result['role']}: assigned FTE exceeds project demand.")
        if not role_result["skills"]:
            alerts.append(
                f"{role_result['role']}: no required skills; no skill-based ranking is available."
            )
    return {
        "roles": result,
        "alerts": alerts,
        "formula": "70 × weighted((level−1)/4) + 20 × weighted(level≥target) + 10 × min(available FTE / remaining assignment slot, 1). Unknown proficiency contributes 0; unknown FTE contributes 0 capacity points and is provisional. Primary/preferred roles never affect score or eligibility.",
        "scale": {1: "Tidak ada pengalaman", 2: "Dasar", 3: "Menengah", 4: "Mahir", 5: "Ahli"},
    }


def export_plan(body: StaffingRequest) -> bytes:
    """Export the recomputed ranking, requirements, and human selections."""
    report = recommend(body)
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = {
        "Assignments": [
            [
                "Project",
                "Role",
                "Talent",
                "Assigned FTE",
                "Current FTE",
                "Capacity FTE",
                "Total Planned FTE",
                "Reason",
                "Status",
            ]
        ],
        "Recommendations": [
            [
                "Role",
                "Rank",
                "Talent",
                "Primary Role",
                "Preferred Roles",
                "Score",
                "Skill Points / 70",
                "Coverage Points / 20",
                "Capacity Points / 10",
                "Available FTE",
                "Provisional",
            ]
        ],
        "Requirements": [
            ["WBS ID", "Item", "Role", "Pillar", "Skill", "UCP Class", "Weight", "Target", "Reason"]
        ],
        "Skill Scores": [
            ["Role", "Talent", "Pillar", "Skill", "Level", "Target", "Weight", "Meets Target"]
        ],
        "Method": [
            ["Item", "Value"],
            ["Project", body.project.name],
            ["Source", body.project.source],
            ["FTE method", body.project.method],
            ["Formula", report["formula"]],
            [
                "Capacity basis",
                "Current FTE must cover the same project period. All proposed role allocations are additive/concurrent. Monthly phasing is not inferred.",
            ],
            [
                "Decision",
                "Human-selected planning proposal. Ranking does not automatically approve assignment.",
            ],
            ["Scale", "1 Tidak ada pengalaman; 2 Dasar; 3 Menengah; 4 Mahir; 5 Ahli"],
            *[["Alert", a] for a in report["alerts"]],
            *[["Import warning", a] for a in body.project.warnings],
        ],
    }
    by_id = {t.id: t for t in body.talents}
    for a in body.assignments:
        t = by_id[a.talent_id]
        total = (
            None
            if t.current_fte is None
            else t.current_fte + sum(x.fte for x in body.assignments if x.talent_id == t.id)
        )
        status = (
            "Provisional"
            if total is None
            else "Over capacity"
            if total > t.capacity_fte + 1e-9
            else "Human selected"
        )
        sheets["Assignments"].append(
            [
                body.project.name,
                a.role,
                t.name,
                a.fte,
                t.current_fte,
                t.capacity_fte,
                total,
                a.reason,
                status,
            ]
        )
    for role in report["roles"]:
        for rank, c in enumerate(role["candidates"], 1):
            sheets["Recommendations"].append(
                [
                    role["role"],
                    rank,
                    c["name"],
                    c["primary_role"],
                    c["preferred_roles"],
                    c["score"],
                    c["skill_score"],
                    c["coverage_score"],
                    c["capacity_score"],
                    c["available_fte"],
                    c["provisional"],
                ]
            )
            for detail in c["details"]:
                sheets["Skill Scores"].append(
                    [
                        role["role"],
                        c["name"],
                        detail["pillar"],
                        detail["skill"],
                        detail["level"],
                        detail["target"],
                        detail["weight"],
                        detail["meets_target"],
                    ]
                )
    for r in body.project.requirements:
        sheets["Requirements"].append(
            [
                r.item_id,
                r.item_name,
                r.role,
                r.pillar,
                r.skill,
                r.classification,
                r.weight,
                r.target,
                r.reason,
            ]
        )
    from openpyxl.styles import Alignment, Font, PatternFill

    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            # Treat arbitrary user text as text, including leading '='.
            sheet.append(
                [
                    "'" + v if isinstance(v, str) and v.startswith(("=", "+", "-", "@")) else v
                    for v in row
                ]
            )
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24496C")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cells in sheet.columns:
            letter = cells[0].column_letter
            sheet.column_dimensions[letter].width = min(
                70, max(15, max(len(str(c.value or "")) for c in cells) + 2)
            )
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
