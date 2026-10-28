"""Scoring, cross-role eligibility, capacity, imports, and assignment exports."""

import io

import pytest
from api.index import app
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from talent_estimation.catalog import catalog_to_excel, load_catalog_file, load_default_catalog
from talent_estimation.staffing import (
    Assignment,
    Demand,
    Requirement,
    StaffingRequest,
    Talent,
    TalentProject,
    export_plan,
    project_from_run,
    read_project,
    read_talents,
    recommend,
    skill_key,
)


def body(level=5, fte=0):
    req = Requirement(
        item_id="1", role="DA", pillar="Querying & Transformation", skill="Advanced SQL Execution"
    )
    t = Talent(
        id="a",
        name="Engineer A",
        primary_role="Data Engineer",
        preferred_roles="Data Engineer",
        current_fte=fte,
        skills={skill_key(req.role, req.pillar, req.skill): level},
    )
    return StaffingRequest(
        project=TalentProject(requirements=[req], demands=[Demand(role="DA", fte=1)]), talents=[t]
    )


def test_cross_role_candidate_wins_without_title_or_preference_bias():
    b = body()
    b.talents.append(
        b.talents[0].model_copy(
            update={
                "id": "b",
                "name": "Analyst B",
                "primary_role": "Data Analyst",
                "preferred_roles": "Data Analyst",
                "skills": {next(iter(b.talents[0].skills)): 2},
            }
        )
    )
    candidates = recommend(b)["roles"][0]["candidates"]
    assert candidates[0]["name"] == "Engineer A"
    assert candidates[0]["score"] == 100
    before = candidates[0]["score"]
    b.talents[0].primary_role = "Unrelated role"
    b.talents[0].preferred_roles = "Other"
    assert recommend(b)["roles"][0]["candidates"][0]["score"] == before


def test_level_one_is_no_experience_and_unknown_fte_is_not_zero():
    known = recommend(body(level=1))["roles"][0]["candidates"][0]
    unknown = recommend(body(level=1, fte=None))["roles"][0]["candidates"][0]
    assert known["skill_score"] == 0 and known["coverage_score"] == 0
    assert known["score"] == 10
    assert unknown["available_fte"] is None and unknown["capacity_score"] == 0
    assert unknown["provisional"] and unknown["score"] == 0


def test_missing_skill_is_not_removed_from_denominator():
    b = body()
    b.project.requirements.append(
        Requirement(item_id="1", role="DA", pillar="Other", skill="Not assessed", weight=10)
    )
    c = recommend(b)["roles"][0]["candidates"][0]
    assert c["skill_score"] == 35 and c["coverage_score"] == 10
    assert c["score"] == 55
    assert c["details"][1]["level"] is None


def test_ucp_weighted_proficiency_and_target_coverage():
    b = body(level=5)
    r = Requirement(
        item_id="2",
        role="DA",
        pillar="Other",
        skill="Hard skill",
        classification="Complex",
        weight=30,
        target=4,
    )
    b.project.requirements.append(r)
    b.talents[0].skills[skill_key(r.role, r.pillar, r.skill)] = 1
    c = recommend(b)["roles"][0]["candidates"][0]
    assert c["skill_score"] == 17.5 and c["coverage_score"] == 5
    assert c["score"] == 32.5


def test_concurrent_assignments_share_capacity_across_roles():
    b = body(fte=0.4)
    b.project.demands.append(Demand(role="DE", fte=1))
    b.assignments = [
        Assignment(talent_id="a", role="DA", fte=0.4),
        Assignment(talent_id="a", role="DE", fte=0.3),
    ]
    report = recommend(b)
    assert report["roles"][0]["candidates"][0]["available_fte"] == 0
    assert any("exceeds capacity" in a for a in report["alerts"])


def test_empty_requirements_do_not_create_a_positive_rank():
    b = body()
    b.project.requirements = []
    report = recommend(b)
    assert report["roles"][0]["candidates"][0]["score"] is None
    assert any("no required skills" in a for a in report["alerts"])


def test_duplicate_requirements_are_rejected():
    b = body()
    b.project.requirements.append(b.project.requirements[0])
    with pytest.raises(ValueError, match="Duplicate requirement"):
        recommend(b)


def test_duplicate_talent_and_unknown_assignments_are_rejected():
    b = body()
    with pytest.raises(ValidationError, match="unique"):
        StaffingRequest(project=b.project, talents=b.talents * 2)
    with pytest.raises(ValidationError, match="unknown talent"):
        StaffingRequest(
            project=b.project,
            talents=b.talents,
            assignments=[Assignment(talent_id="bad", role="DA", fte=0.5)],
        )


def xlsx(workbook):
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_roster_fte_and_scores_import_and_strict_score_validation():
    w = Workbook()
    s = w.active
    s.title = "Talent Capability Breakdown"
    s.append(
        ["Nama", "Role Utama", "Preferred Role", "Tribe / Project", "Data Analyst", "Current FTE"]
    )
    s.append([None, None, None, None, "Advanced SQL Execution", None])
    s.append(["Engineer A", "Data Engineer", "Data Engineer", "Other", 5, 0.4])
    parsed = read_talents(xlsx(w), load_default_catalog())
    assert parsed["talents"][0]["current_fte"] == 0.4
    assert next(iter(parsed["talents"][0]["skills"].values())) == 5
    s["E3"] = 2.5
    with pytest.raises(ValueError, match="1–5"):
        read_talents(xlsx(w), load_default_catalog())


@pytest.mark.parametrize("second_header", [False, True])
def test_fte_in_column_e_is_workload_not_a_skill(second_header):
    w = Workbook()
    s = w.active
    s.title = "Talent Capability Breakdown"
    s.append(
        [
            "Nama",
            "Role Utama",
            "Preferred Role",
            "Tribe / Project",
            None if second_header else "FTE",
            "Data Analyst",
        ]
    )
    s.append([None, None, None, None, "FTE" if second_header else None, "Advanced SQL Execution"])
    s.append(["Engineer A", "Data Engineer", "Data Engineer", "Synthetic", 0.4, 5])
    s.append(["Engineer B", "Data Engineer", "Data Engineer", "Synthetic", 0, 3])
    s.append(["Engineer C", "Data Engineer", "Data Engineer", "Synthetic", None, 2])
    client = TestClient(app)
    response = client.post("/api/py/talent/parse", files={"file": ("talents.xlsx", xlsx(w))})
    assert response.status_code == 200, response.text
    parsed = response.json()
    assert [t["current_fte"] for t in parsed["talents"]] == [0.4, 0, None]
    assert parsed["matched_skills"] == 1
    assert [next(iter(t["skills"].values())) for t in parsed["talents"]] == [5, 3, 2]
    assert len(parsed["warnings"]) == 1
    assert parsed["warnings"][0].startswith("1 talents have no Current FTE")
    b = body()
    b.talents = [Talent.model_validate(t) for t in parsed["talents"]]
    candidates = recommend(b)["roles"][0]["candidates"]
    engineer = next(c for c in candidates if c["name"] == "Engineer A")
    assert engineer["available_fte"] == pytest.approx(0.6)
    assert not engineer["provisional"]


def test_validation_import_stops_at_next_summary_section():
    w = Workbook()
    s = w.active
    s.title = "Capability"
    s.append(["WBS ID", "Item WBS", "Role", "Pilar", "Skill", "Bagian pekerjaan"])
    s.append(["1", "Query", "DA", "Querying & Transformation", "Advanced SQL Execution", "SQL"])
    s = w.create_sheet("Breakdown_WBS")
    s.append(["WBS ID", "Role", "Klasifikasi UCP", "Qty", "Peran", "Bobot efektif UCP"])
    s.append(["1", "DA", "Complex", 2, "Utama", 30])
    s = w.create_sheet("Ringkasan")
    s.append(["Kode", "Bottom-up FTE puncak"])
    s.append(["DA", 0.8])
    s.append([])
    s.append(["Kode", "Other section"])
    s.append(["DA", 999])
    p = read_project(xlsx(w))
    assert len(p.demands) == 1 and p.demands[0].fte == 0
    assert p.demands[0].estimates["bottom_up_peak"] == 0.8
    assert p.requirements[0].weight == 30


@pytest.mark.parametrize("requested,expected", [(2, 2), ("0,5", 0.5), (0, 0), (None, 0)])
def test_requested_fte_is_separate_from_estimate(requested, expected):
    w = Workbook()
    s = w.active
    s.title = "Capability"
    s.append(["WBS ID", "Role", "Pilar", "Skill"])
    s.append(["1", "DA", "Querying & Transformation", "Advanced SQL Execution"])
    s = w.create_sheet("Breakdown_WBS")
    s.append(["WBS ID", "Role", "Klasifikasi UCP"])
    s.append(["1", "DA", "Average"])
    s = w.create_sheet("Ringkasan")
    s.append(["Kode", "Role", "Diminta (FTE)", "Bottom-up FTE puncak", "UCP FTE rata-rata"])
    s.append(["DA", "Data Analyst", requested, 1.23, 0.8])
    s.append(["UX", "UX", 0, 0, 0])
    p = read_project(xlsx(w))
    assert p.demands[0].fte == expected
    assert p.demands[0].estimates["bottom_up_peak"] == 1.23
    assert p.demands[0].estimates["ucp_average"] == 0.8
    assert p.demands[1].fte == 0


def test_demand_alerts_only_for_positive_unselected_roles_and_capacity_still_warns():
    b = body(fte=0.5)
    b.project.demands.extend([Demand(role="UX", fte=0), Demand(role="UI")])
    report = recommend(b)
    assert len(report["alerts"]) == 1
    assert "belum ada talent dipilih" in report["alerts"][0]
    b.assignments = [Assignment(talent_id="a", role="DA", fte=0.2)]
    report = recommend(b)
    assert report["roles"][0]["remaining_fte"] == 0.8
    assert report["alerts"] == []
    b.assignments[0].fte = 1
    assert any("exceeds capacity" in a for a in recommend(b)["alerts"])


def test_legacy_engineer_labels_import_as_ai_engineer():
    t = Talent(
        id="a",
        name="A",
        primary_role="Engineer",
        preferred_roles="Data Engineer, Engineer, AI Engineer",
    )
    assert t.primary_role == "AI Engineer"
    assert t.preferred_roles == "Data Engineer, AI Engineer, AI Engineer"
    assert skill_key("Engineer", "P", "S") == skill_key("AI Engineer", "P", "S")


async def test_validation_estimates_without_cached_formulas_match_current_run(demo_inputs, catalog):
    from talent_estimation.agents import run_pipeline
    from talent_estimation.domain.models import RunResult
    from talent_estimation.excel import build_workbook

    async for event in run_pipeline(demo_inputs, None):
        if event["type"] == "result":
            run = RunResult.model_validate(event["run"])
    run.role_inputs[0].existing_fte = 0.25
    content = build_workbook(run, catalog)
    restored = read_project(content)
    direct = project_from_run(run)
    assert restored.estimate_note
    assert not restored.warnings
    assert len(restored.demands) == len(direct.demands)
    for imported, current in zip(restored.demands, direct.demands, strict=True):
        assert imported.role == current.role
        assert imported.fte == current.fte
        assert imported.existing_fte == current.existing_fte
        assert imported.estimates == pytest.approx(current.estimates)
        if current.divergence is None:
            assert imported.divergence is None
        else:
            assert imported.divergence == pytest.approx(current.divergence)
    # Saved numeric results remain authoritative when only other formulas lack caches.
    w = load_workbook(io.BytesIO(content))
    w["Ringkasan"]["G7"] = 123
    assert read_project(xlsx(w)).demands[0].estimates["ucp_average"] == 123


@pytest.mark.parametrize(
    "fte_value,expected",
    [(0.5, 0.5), ("0,5", 0.5), (" 0,50 ", 0.5), ("0.5", 0.5), ("0,00", 0), ("50%", 0.5)],
)
def test_locale_fte_values_and_excel_number_format(fte_value, expected):
    w = Workbook()
    s = w.active
    s.title = "Talent Capability Breakdown"
    s.append(["Nama", "Role Utama", "Preferred Role", "Tribe / Project", "FTE", "Data Analyst"])
    s.append([None, None, None, None, None, "Advanced SQL Execution"])
    s.append(["Engineer A", "Data Engineer", "Data Engineer", "Synthetic", fte_value, 5])
    s["E3"].number_format = "#,##0.00"
    parsed = read_talents(xlsx(w), load_default_catalog())
    assert parsed["talents"][0]["current_fte"] == expected
    assert parsed["fte_import"] == {
        "columns": ["E"],
        "loaded": 1,
        "missing": 0,
        "formulas_without_cached_value": 0,
    }


def test_uncached_fte_formula_has_an_actionable_warning():
    w = Workbook()
    s = w.active
    s.title = "Talent Capability Breakdown"
    s.append(["Nama", "Role Utama", "Preferred Role", "Tribe / Project", "FTE", "Data Analyst"])
    s.append([None, None, None, None, None, "Advanced SQL Execution"])
    s.append(["Engineer A", "Data Engineer", "Data Engineer", "Synthetic", "=0.2+0.3", 5])
    parsed = read_talents(xlsx(w), load_default_catalog())
    assert parsed["talents"][0]["current_fte"] is None
    assert parsed["fte_import"]["formulas_without_cached_value"] == 1
    assert any("E3" in w and "Recalculate and save" in w for w in parsed["warnings"])


def test_missing_fte_header_is_reported_in_import_summary():
    w = Workbook()
    s = w.active
    s.title = "Talent Capability Breakdown"
    s.append(["Nama", "Data Analyst"])
    s.append([None, "Advanced SQL Execution"])
    s.append(["Engineer A", 5])
    parsed = read_talents(xlsx(w), load_default_catalog())
    assert parsed["fte_import"]["columns"] == []
    assert any("No FTE or Current FTE column" in w for w in parsed["warnings"])


@pytest.mark.parametrize("top,label", [("Current", "FTE"), ("Data Engineer", "FTE")])
def test_fte_header_recognized_below_grouped_or_split_header(top, label):
    w = Workbook()
    s = w.active
    s.title = "Talent Capability Breakdown"
    s.append(["Nama", None, None, None, top, "Data Analyst"])
    s.append([None, None, None, None, label, "Advanced SQL Execution"])
    s.append(["Engineer A", None, None, None, "0,5", 5])
    parsed = read_talents(xlsx(w), load_default_catalog())
    assert parsed["talents"][0]["current_fte"] == 0.5


def test_invalid_decimal_fte_is_not_silently_treated_as_unknown():
    w = Workbook()
    s = w.active
    s.title = "Talent Capability Breakdown"
    s.append(["Nama", "FTE", "Data Analyst"])
    s.append([None, None, "Advanced SQL Execution"])
    s.append(["Engineer A", "0,x", 5])
    with pytest.raises(ValueError, match="B3: invalid FTE"):
        read_talents(xlsx(w), load_default_catalog())


def test_catalog_overlay_and_round_trip_definitions():
    c = load_default_catalog()
    assert (
        sum(
            len(skills)
            for code in ["DE", "DA", "DS", "AI"]
            for skills in c.roles[code].pillars.values()
        )
        == 127
    )
    assert len(c.roles["DE"].skill_definitions) == 26
    restored = load_catalog_file("catalog.xlsx", catalog_to_excel(c))
    assert restored.roles["DA"].pillars == c.roles["DA"].pillars
    assert restored.roles["DS"].skill_definitions == c.roles["DS"].skill_definitions


def test_excel_includes_human_selection_and_prevents_formula_injection():
    b = body()
    b.talents[0].name = "=2+2"
    b.assignments = [Assignment(talent_id="a", role="DA", fte=0.5, reason="Human choice")]
    w = load_workbook(io.BytesIO(export_plan(b)))
    assert w["Assignments"]["C2"].data_type != "f"
    assert w["Assignments"]["D2"].value == 0.5
    assert w["Assignments"]["H2"].value == "Human choice"
    assert len(w.sheetnames) == 5


def test_api_ranking_and_export_routes_and_invalid_inputs():
    client = TestClient(app)
    b = body()
    response = client.post("/api/py/talent/recommend", json=b.model_dump())
    assert response.status_code == 200
    assert response.json()["roles"][0]["candidates"][0]["score"] == 100
    assert client.post("/api/py/talent/export", json=b.model_dump()).status_code == 200
    bad = b.model_dump()
    bad["talents"][0]["skills"] = {"test": 6}
    assert client.post("/api/py/talent/recommend", json=bad).status_code == 422
