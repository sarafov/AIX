"""Editable, source-grounded resource SLA drafts and document exports."""

from __future__ import annotations

import io
import re
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from pydantic import BaseModel, Field

from talent_estimation.catalog import Catalog
from talent_estimation.staffing import ROLE_NAMES, StaffingRequest, norm, role_code

PLACEHOLDER = "[Diisi reviewer]"
INTAKE = """Backlog Lock-in & No Mid-Sprint Injections
Scope pekerjaan dikunci saat Sprint Planning. Pekerjaan baru di tengah sprint dimasukkan ke backlog sprint berikutnya.

Effort Estimation Alignment
Estimasi bobot pekerjaan disepakati Project Owner dan talent menggunakan konsensus Story Points atau jam kerja efektif sebelum sprint dimulai.

Capacity Prorating & Spillover Allocation
Cuti resmi dan hari libur diperhitungkan secara prorata. Spillover sprint sebelumnya diprioritaskan sebelum pekerjaan baru.

Talent Request & Penambahan Kapasitas
Penambahan talent/FTE di luar tabel memerlukan permintaan resmi dari Tribe/Project Owner melalui TMS dan/atau Nota Dinas serta pengesahan Chapter terkait."""
COMMUNICATION = """Asynchronous-First Standard
Dokumentasi, pembaruan tugas, dan diskusi teknis dipusatkan melalui Taiga, Activity Logger/Diarium, Microsoft Teams, serta Nota Dinas/TMS untuk penugasan resmi.

Response Time SLA
Pertanyaan/diskusi normal: direspons dalam 4-6 jam kerja. Critical Blocker: direspons maksimal 2 jam kerja.

Agile Ritual Participation
Kehadiran sinkron wajib pada Sprint Planning dan Sprint Review/Demo. Daily Standup dilakukan secara asinkron dengan format: output yang diselesaikan kemarin, rencana kerja hari ini, dan kendala/blocker."""
DOD = """Code Review & Approval: kode telah melalui peer review dan persetujuan pull request.
Quality Verification: lulus unit testing dan integration testing sesuai standar coverage yang disepakati.
Handover & Documentation: dokumentasi teknis, API specification, dan panduan operasional telah diperbarui dan diserahterimakan.
Anti Scope-Creep Guarantee: acceptance criteria tidak diubah di tengah pelaksanaan tanpa persetujuan ulang para pihak.
Data/Design Documentation: dokumentasi data dan desain disusun serta dipelihara menggunakan template Chapter terkait.
Use Case Acceptance: deliverable lulus UAT dan sign-off sesuai jadwal serta lingkungan yang disepakati dalam dokumen proyek."""
ESCALATION = """Buffer & Spillover Handling: potensi keterlambatan dan penyesuaian jadwal/buffer dibahas secara terbuka.
Escalation Threshold: keterlambatan deliverable lebih dari 20% selama 2 sprint berturut-turut memerlukan Capacity Review formal antara Project Owner dan Chapter Lead.
Shared Talent Contention: benturan penugasan diselesaikan melalui diskusi Project Owner, PM, dan Chapter Lead. Alokasi proyek tidak melebihi FTE pada Bagian 2."""


class SlaRow(BaseModel):
    """One editable assignment and its deliverable boundaries."""

    role: str = Field(default="", max_length=200)
    talent: str = Field(default="", max_length=300)
    fte: float = Field(default=0, ge=0, le=10, allow_inf_nan=False)
    in_scope: str = Field(default="", max_length=12000)
    out_scope: str = Field(default="", max_length=12000)


class SlaSeed(BaseModel):
    """Human choices from the current plan or an assignment export."""

    project: str = ""
    rows: list[SlaRow] = Field(default_factory=list, max_length=500)
    tasks: dict[str, list[str]] = Field(default_factory=dict)
    warnings: list[str] = Field(default_factory=list)


class SlaDraft(BaseModel):
    """The complete editable content, independent of either export format."""

    title: str = "Deliverable-Based Resource SLA & Commitment Agreement"
    project_name: str = ""
    chapter: str = PLACEHOLDER
    period: str = PLACEHOLDER
    owner_pm: str = PLACEHOLDER
    background: str = PLACEHOLDER
    objectives: str = PLACEHOLDER
    engagement: str = "Talent tetap berada di bawah Chapter asal dan dialokasikan melalui Matrix Sharing Model. Komitmen berfokus pada output, deliverable, Definition of Done, dan kualitas keluaran yang disepakati pada setiap sprint."
    assignment_terms: str = PLACEHOLDER
    deliverable_notes: str = PLACEHOLDER
    rows: list[SlaRow] = Field(default_factory=list, max_length=500)
    intake: str = INTAKE
    communication: str = COMMUNICATION
    definition_of_done: str = DOD
    escalation: str = ESCALATION
    sources: list[str] = Field(default_factory=list, max_length=30)
    warnings: list[str] = Field(default_factory=list, max_length=100)


class SlaDraftRequest(BaseModel):
    """Draft inputs without server-side persistence."""

    seed: SlaSeed | None = None
    staffing: StaffingRequest | None = None
    project_info: dict[str, Any] = Field(default_factory=dict)
    catalog: Catalog | None = None


def seed_from_staffing(body: StaffingRequest) -> SlaSeed:
    """Use selected assignments only, never automatically assign recommended talents."""
    by_id = {t.id: t for t in body.talents}
    tasks: dict[str, list[str]] = {}
    for r in body.project.requirements:
        task = r.item_name or r.skill
        if task not in tasks.setdefault(r.role, []):
            tasks[r.role].append(task)
    return SlaSeed(
        project=body.project.name,
        tasks=tasks,
        rows=[
            SlaRow(role=a.role, talent=by_id[a.talent_id].name, fte=a.fte) for a in body.assignments
        ],
    )


def read_assignment_plan(content: bytes) -> SlaSeed:
    """Read the exported Assignments table without importing recommendations as choices."""
    w = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    if any((s.max_row or 0) > 20000 or (s.max_column or 0) > 200 for s in w):
        raise ValueError("Workbook terlalu besar.")
    if "Assignments" not in w.sheetnames:
        raise ValueError("Sheet Assignments tidak ditemukan. Unggah Talent_Assignment_Plan.xlsx.")
    records = list(w["Assignments"].values)
    headers = [norm(v) for v in records[0]]
    if not {"role", "talent", "assigned fte"} <= set(headers):
        raise ValueError("Kolom Role, Talent, dan Assigned FTE wajib tersedia.")
    seed = SlaSeed()
    for values in records[1:]:
        row = dict(zip(headers, values, strict=True))
        if not row.get("talent"):
            continue
        if not row.get("role"):
            raise ValueError("Role talent kosong.")
        seed.project = seed.project or str(row.get("project") or "")
        seed.rows.append(
            SlaRow(
                role=role_code(row["role"]),
                talent=str(row["talent"]),
                fte=float(str(row["assigned fte"]).replace(",", ".")),
            )
        )
        if row.get("status") in {"Over capacity", "Provisional"}:
            seed.warnings.append(f"{row['talent']}: {row['status']} pada assignment plan.")
    if not seed.rows:
        raise ValueError("Belum ada talent dipilih pada Assignments.")
    if "Requirements" in w.sheetnames:
        rows = list(w["Requirements"].values)
        h = [norm(v) for v in rows[0]]
        for values in rows[1:]:
            row = dict(zip(h, values, strict=True))
            role = role_code(row.get("role"))
            task = str(row.get("item") or row.get("skill") or "")
            if task and task not in seed.tasks.setdefault(role, []):
                seed.tasks[role].append(task)
    return seed


def _excerpt(text: str, labels: list[str], limit: int = 1800) -> str:
    lines = [
        line.strip() for line in text.splitlines() if line.strip() and not line.strip().isdigit()
    ]
    for label in labels:
        for i, line in enumerate(lines):
            heading = re.sub(r"^(?:[a-z]|\d+)\s+", "", norm(line))
            if heading == label:
                block: list[str] = []
                for following in lines[i + 1 :]:
                    if block and re.match(r"^(?:[A-Z]\.|\d+\.)\s+[A-Z]", following):
                        break
                    block.append(following)
                    if len(" ".join(block)) >= limit:
                        break
                if block:
                    excerpt = " ".join(block)
                    if len(excerpt) > limit:
                        end = excerpt.rfind(". ", 0, limit)
                        return (
                            excerpt[: end + 1]
                            if end > limit // 2
                            else excerpt[:limit].rsplit(" ", 1)[0] + "…"
                        )
                    return excerpt
    return ""


def make_sla(body: SlaDraftRequest, documents: list[tuple[str, str]], catalog: Catalog) -> SlaDraft:
    """Assemble extractive draft content with explicit missing facts and provenance."""
    seed = seed_from_staffing(body.staffing) if body.staffing else body.seed
    if not seed or not seed.rows:
        raise ValueError(
            "Pilih talent pada rencana penugasan atau unggah Talent Assignment Plan terlebih dahulu."
        )
    text = "\n\n".join(t for _, t in documents)
    context = str(body.project_info.get("context") or "").strip()
    if len(text.strip()) < 40 and len(context) < 40:
        raise ValueError(
            "Informasi proyek belum tersedia. Unggah KAK/project charter/dokumen pendukung, atau isi konteks proyek minimal 40 karakter."
        )
    name = str(body.project_info.get("name") or seed.project)
    if not name or norm(name) in {"project", "imported validation"}:
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        for i, line in enumerate(lines[:20]):
            if "kerangka acuan kerja" in norm(line) or "project charter" in norm(line):
                name = " ".join(lines[i + 1 : i + 3])
                break
    draft = SlaDraft(
        project_name=name or PLACEHOLDER, sources=[n for n, _ in documents], warnings=seed.warnings
    )
    if context:
        draft.sources.append("Informasi proyek dari hasil estimasi saat ini")
    draft.background = (
        _excerpt(text, ["gambaran umum", "project context", "background", "latar belakang"], 2400)
        or context
        or PLACEHOLDER
    )
    draft.objectives = (
        _excerpt(
            text,
            ["maksud dan tujuan", "tujuan dan sasaran", "tujuan kegiatan", "objectives", "tujuan"],
            1800,
        )
        or PLACEHOLDER
    )
    draft.deliverable_notes = (
        _excerpt(
            text,
            [
                "catatan penugasan",
                "jangka waktu",
                "waktu pencapaian keluaran",
                "waktu pelaksanaan",
                "periode pelaksanaan",
                "delivery notes",
            ],
            1200,
        )
        or PLACEHOLDER
    )
    for field, label in [
        ("chapter", r"(?:Talent Source\s*/\s*Chapter|Chapter)"),
        ("owner_pm", r"(?:Project Owner\s*/?\s*PM|Pemilik Proyek|Penanggung Jawab)"),
    ]:
        match = re.search(label + r"\s*:\s*([^\n]{3,220})", text, flags=re.I)
        if match:
            setattr(draft, field, match[1].strip())
    start, end = body.project_info.get("start"), body.project_info.get("end")
    draft.period = (
        f"{start} s.d. {end}"
        if start and end
        else _excerpt(
            text,
            [
                "jangka waktu pelaksanaan",
                "waktu pencapaian keluaran",
                "waktu pelaksanaan",
                "periode pelaksanaan",
            ],
            700,
        )
        or PLACEHOLDER
    )
    if not (start and end) and draft.period != PLACEHOLDER:
        months = re.findall(
            r"(?:Januari|Februari|Maret|April|Mei|Juni|Juli|Agustus|September|Oktober|November|Desember)\s+\d{4}",
            draft.period,
            re.I,
        )
        if len(months) >= 2:
            draft.period = (
                ("Rencana: " if "direncanakan" in draft.period.lower() else "")
                + months[0]
                + " s.d. "
                + months[-1]
            )
    draft.assignment_terms = (
        _excerpt(text, ["ketentuan penugasan", "engagement model", "assignment terms"], 1000)
        or PLACEHOLDER
    )
    for selected in seed.rows:
        code = role_code(selected.role)
        profile = catalog.roles.get(code)
        tasks = seed.tasks.get(code, [])
        scope = profile.in_scope if profile else PLACEHOLDER
        if tasks:
            scope += "\nDeliverable proyek untuk ditinjau: " + "; ".join(tasks[:12])
        scope += "\nKPI / acceptance criteria: " + PLACEHOLDER
        draft.rows.append(
            selected.model_copy(
                update={
                    "role": ROLE_NAMES.get(code, profile.name if profile else selected.role),
                    "in_scope": selected.in_scope or scope,
                    "out_scope": selected.out_scope
                    or (profile.out_of_scope if profile else PLACEHOLDER),
                }
            )
        )
    missing = [
        label
        for field, label in [
            ("project_name", "nama proyek"),
            ("chapter", "Chapter"),
            ("period", "periode"),
            ("owner_pm", "Project Owner/PM"),
            ("background", "latar belakang"),
            ("objectives", "objektif"),
            ("assignment_terms", "ketentuan penugasan"),
            ("deliverable_notes", "catatan deliverable"),
        ]
        if getattr(draft, field) == PLACEHOLDER
    ]
    if missing:
        draft.warnings.append("Lengkapi dari dokumen atau reviewer: " + ", ".join(missing) + ".")
    draft.warnings.append(
        "Bagian 3-6 mengikuti contoh SLA sebagai ketentuan draf yang dapat diubah; bukan persetujuan penugasan."
    )
    return draft


def _sections(d: SlaDraft) -> list[tuple[str, str]]:
    return [
        (
            "1. Project Context & Engagement Model",
            "Latar belakang proyek\n"
            + d.background
            + "\n\nObjektif strategis\n"
            + d.objectives
            + "\n\nModel keterlibatan\n"
            + d.engagement
            + "\n\nKetentuan penugasan\n"
            + d.assignment_terms,
        ),
        ("2. FTE & KPI Deliverable", d.deliverable_notes),
        ("3. Work Intake & Sprint Planning Rules", d.intake),
        ("4. Communication SLA & Async Operating Rhythm", d.communication),
        ("5. Definition of Done & Sign-Off Criteria", d.definition_of_done),
        ("6. Escalation Protocol", d.escalation),
    ]


def export_sla_word(d: SlaDraft) -> bytes:
    """Export editable Word paragraphs and a native table from the exact draft state."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.65)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Arial", Pt(10)
    style.paragraph_format.space_after = Pt(6)
    for name in ["Title", "Heading 1"]:
        doc.styles[name].font.name = "Arial"
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    title_properties = doc.styles["Title"].element.get_or_add_pPr()
    for border in list(title_properties.findall(qn("w:pBdr"))):
        title_properties.remove(border)
    doc.styles["Title"].font.size = Pt(20)
    doc.add_heading(d.title, 0)
    doc.add_paragraph("DRAF - untuk ditinjau")
    for label, value in [
        ("Project Name", d.project_name),
        ("Talent Source / Chapter", d.chapter),
        ("Cycle / Period", d.period),
        ("Project Owner / PM", d.owner_pm),
    ]:
        p = doc.add_paragraph()
        p.add_run(label + ": ").bold = True
        p.add_run(value)
    for index, (heading, text) in enumerate(_sections(d)):
        doc.add_heading(heading, 1)
        for paragraph in text.split("\n\n"):
            doc.add_paragraph(paragraph)
        if index == 1:
            table = doc.add_table(rows=1, cols=6)
            table.style = "Light Shading Accent 1"
            headers = [
                "No",
                "Role / Skillstream",
                "Nama Talent",
                "FTE",
                "In-Scope Deliverables",
                "Out-of-Scope",
            ]
            for cell, value in zip(table.rows[0].cells, headers, strict=True):
                cell.text = value
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            table.autofit = False
            widths = [0.3, 0.85, 1.0, 0.4, 2.35, 2.3]
            for column, width in zip(table.columns, widths, strict=True):
                column.width = Inches(width)
            for cell, width in zip(table.rows[0].cells, widths, strict=True):
                cell.width = Inches(width)
            for i, row in enumerate(d.rows, 1):
                cells = table.add_row().cells
                for cell, value, width in zip(
                    cells,
                    [str(i), row.role, row.talent, f"{row.fte:g}", row.in_scope, row.out_scope],
                    widths,
                    strict=True,
                ):
                    cell.text, cell.width = value, Inches(width)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(4)
                        for run in p.runs:
                            run.font.size = Pt(9)
            doc.add_paragraph(f"Total alokasi: {sum(r.fte for r in d.rows):g} FTE")
    doc.add_heading("Sumber dan catatan review", 1)
    for item in d.sources + d.warnings:
        doc.add_paragraph(item)
    footer = section.footer.paragraphs[0]
    footer.add_run("Draf Resource SLA | ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def export_sla_pdf(d: SlaDraft) -> bytes:
    """Render the edited draft into a paginated PDF with repeated table headings."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    output = io.BytesIO()
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="BodySla",
            fontName="Helvetica",
            fontSize=10,
            leading=13,
            spaceAfter=7,
            alignment=TA_LEFT,
        )
    )
    styles.add(ParagraphStyle(name="CellSla", fontName="Helvetica", fontSize=8, leading=11))

    def para(text: str, style: str = "BodySla") -> Any:
        return Paragraph(escape(text).replace("\n", "<br/>"), styles[style])

    story = [para(d.title, "Title"), para("DRAF - untuk ditinjau")]
    for label, value in [
        ("Project Name", d.project_name),
        ("Talent Source / Chapter", d.chapter),
        ("Cycle / Period", d.period),
        ("Project Owner / PM", d.owner_pm),
    ]:
        story.append(para(label + ": " + value))
    for index, (heading, text) in enumerate(_sections(d)):
        story.append(para(heading, "Heading2"))
        story.extend(para(p) for p in text.split("\n\n"))
        if index == 1:
            data = [
                [
                    para(v, "CellSla")
                    for v in [
                        "No",
                        "Role / Skillstream",
                        "Nama Talent",
                        "FTE",
                        "In-Scope Deliverables",
                        "Out-of-Scope",
                    ]
                ]
            ]
            for i, row in enumerate(d.rows, 1):
                data.append(
                    [
                        para(v, "CellSla")
                        for v in [
                            str(i),
                            row.role,
                            row.talent,
                            f"{row.fte:g}",
                            row.in_scope,
                            row.out_scope,
                        ]
                    ]
                )
            table = LongTable(
                data, colWidths=[24, 70, 85, 32, 250, 243], repeatRows=1, splitInRow=1
            )
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E9EEF4")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            story.extend(
                [table, Spacer(1, 8), para(f"Total alokasi: {sum(r.fte for r in d.rows):g} FTE")]
            )
    story.append(para("Sumber dan catatan review", "Heading2"))
    story.extend(para(item, "CellSla") for item in d.sources + d.warnings)

    def page_number(canvas: Any, document: Any) -> None:
        canvas.setFont("Helvetica", 8)
        canvas.drawString(44, 24, "Draf Resource SLA")
        canvas.drawRightString(748, 24, str(document.page))

    SimpleDocTemplate(
        output,
        pagesize=landscape(letter),
        rightMargin=44,
        leftMargin=44,
        topMargin=36,
        bottomMargin=40,
    ).build(story, onFirstPage=page_number, onLaterPages=page_number)
    return output.getvalue()
